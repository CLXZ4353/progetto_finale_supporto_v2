# -*- coding: utf-8 -*-
"""Genera la presentazione PPTX della v2 (12 slide).

Legge le metriche reali da eval/results e riusa il diagramma di architettura.
"""
from __future__ import annotations

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
RESULTS = ROOT / "eval" / "results"
OUT = DOCS / "presentazione_v2.pptx"
CHART = DOCS / "valutazione_v2.png"
ARCH = DOCS / "architettura.png"

# --- palette ---------------------------------------------------------------
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BROWN = RGBColor(0x6B, 0x44, 0x29)
BROWN_DARK = RGBColor(0x3D, 0x2C, 0x20)
GREEN = RGBColor(0x4F, 0x6B, 0x4D)
GOLD = RGBColor(0xC8, 0x9B, 0x5A)
INK = RGBColor(0x25, 0x25, 0x25)
GREY = RGBColor(0x5B, 0x5B, 0x5B)
LIGHT_BROWN = RGBColor(0xF3, 0xE9, 0xDF)
LIGHT_GREEN = RGBColor(0xE8, 0xEF, 0xE6)
LIGHT_BLUE = RGBColor(0xE9, 0xEF, 0xF7)
LIGHT_RED = RGBColor(0xF6, 0xE7, 0xE4)
LINE = RGBColor(0xDE, 0xD2, 0xC4)
PALE = RGBColor(0xF7, 0xF3, 0xEE)

FONT = "Calibri"
FONT_MONO = "Consolas"

# --- dati ------------------------------------------------------------------
final = json.loads((RESULTS / "latest.json").read_text(encoding="utf-8"))
fs = final["summary"]

llama = None
for path in sorted(RESULTS.glob("local_2026*.json"), reverse=True):
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        continue
    summary = payload.get("summary", {})
    if "llama-3.2-3b" in str(summary.get("model", "")) and summary.get("total") == 31:
        llama = payload
        break
ls = (llama or {}).get("summary", {})


def fmt_pct(value) -> str:
    if value is None:
        return "n/d"
    rounded = (Decimal(str(value)) * 100).quantize(
        Decimal("0.1"), rounding=ROUND_HALF_UP
    )
    return f"{rounded:.1f}%".replace(".", ",")


def fmt_s(value) -> str:
    if value is None:
        return "n/d"
    return f"{value:.2f} s".replace(".", ",")


# --- grafico di supporto ---------------------------------------------------
if not CHART.exists():
    profiles = [(name, summary, color) for name, summary, color in (
        ("LFM3 1.7B", fs, "#6b4429"),
        ("Llama 3.2 3B Unsloth", ls, "#4f6b4d"),
    ) if summary]
    fig, axes = plt.subplots(1, 2, figsize=(10.6, 2.9))
    names = [name for name, _, _ in profiles]
    colors_bar = [color for _, _, color in profiles]
    rates = [summary["pass_rate"] * 100 for _, summary, _ in profiles]
    latency = [summary["avg_latency_s"] for _, summary, _ in profiles]
    axes[0].bar(names, rates, color=colors_bar, width=0.52)
    axes[0].set_ylim(0, 112)
    axes[0].set_ylabel("pass rate (%)")
    axes[0].set_title("31 scenari superati")
    for index, value in enumerate(rates):
        axes[0].text(index, value + 2, f"{value:.1f}%".replace(".", ","), ha="center", fontsize=9)
    axes[1].bar(names, latency, color=colors_bar, width=0.52)
    top = max(latency) * 1.28
    axes[1].set_ylim(0, top)
    axes[1].set_ylabel("secondi per scenario")
    axes[1].set_title("Latenza media")
    for index, value in enumerate(latency):
        axes[1].text(index, value + top * 0.04, f"{value:.1f} s".replace(".", ","), ha="center", fontsize=9)
    for ax in axes:
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", alpha=0.25)
    fig.subplots_adjust(left=0.08, right=0.98, top=0.86, bottom=0.18, wspace=0.3)
    fig.savefig(CHART, dpi=190, facecolor="white", bbox_inches="tight")
    plt.close(fig)

# --- helpers ---------------------------------------------------------------
TOTAL_SLIDES = 12

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_bg(slide, color: RGBColor = CREAM) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text="",
    *,
    size=14,
    bold=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    italic=False,
    line_spacing=1.0,
    space_after=0,
    font=FONT,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = str(text).split("\n")
    for index, line in enumerate(lines):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after if index < len(lines) - 1 else 0)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
    return box


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    items,
    *,
    size=14,
    color=INK,
    gap=8,
    indent=0.24,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = item
        paragraph.space_after = Pt(gap)
        paragraph.level = 0
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = line_spacing
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color
        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(indent))))
        pPr.set("indent", str(-int(Inches(indent))))
        bu_font = OxmlElement("a:buFont")
        bu_font.set("typeface", "Arial")
        bu_char = OxmlElement("a:buChar")
        bu_char.set("char", "•")
        pPr.append(bu_font)
        pPr.append(bu_char)
    return box


def add_rect(
    slide,
    left,
    top,
    width,
    height,
    *,
    fill=WHITE,
    line=LINE,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    radius=0.08,
):
    obj = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    if line is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = line
        obj.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            obj.adjustments[0] = radius
        except Exception:
            pass
    obj.shadow.inherit = False
    return obj


def add_card(
    slide,
    left,
    top,
    width,
    height,
    title,
    body,
    *,
    accent=GOLD,
    fill=WHITE,
    title_size=18,
    body_size=13,
    body_lines=None,
):
    add_rect(slide, left, top, width, height, fill=fill)
    add_rect(slide, left, top, width, 0.07, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(slide, left + 0.22, top + 0.22, width - 0.44, 0.42, title, size=title_size, bold=True, color=BROWN_DARK)
    if body_lines:
        add_bullets(
            slide,
            left + 0.24,
            top + 0.78,
            width - 0.48,
            height - 0.95,
            body_lines,
            size=body_size,
            gap=6,
        )
    else:
        add_text(
            slide,
            left + 0.22,
            top + 0.78,
            width - 0.44,
            height - 0.95,
            body,
            size=body_size,
            color=GREY,
            line_spacing=1.04,
        )


def add_chip(slide, left, top, width, height, text, *, fill=LIGHT_BROWN, color=BROWN_DARK, size=13):
    add_rect(slide, left, top, width, height, fill=fill, line=None)
    add_text(
        slide,
        left + 0.08,
        top + 0.04,
        width - 0.16,
        height - 0.08,
        text,
        size=size,
        bold=True,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def add_arrow(slide, left, top, width=0.28, *, color=GOLD):
    obj = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(0.22))
    obj.fill.solid()
    obj.fill.fore_color.rgb = color
    obj.line.fill.background()
    obj.shadow.inherit = False


def add_title_block(slide, kicker, title, subtitle=None):
    add_text(slide, 0.6, 0.32, 11.8, 0.26, kicker.upper(), size=10, bold=True, color=GOLD)
    add_text(slide, 0.6, 0.56, 12.1, 0.62, title, size=29, bold=True, color=BROWN_DARK)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.18), Inches(0.7), Inches(0.045))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()
    if subtitle:
        add_text(slide, 0.6, 1.3, 12.0, 0.34, subtitle, size=13, color=GREY)


def add_footer(slide, number):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.02), Inches(12.13), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    add_text(slide, 0.6, 7.1, 9.5, 0.2, "Casa & Regali · assistente clienti · v2", size=8.5, color=GREY)
    add_text(slide, 10.8, 7.1, 1.93, 0.2, f"{number} / {TOTAL_SLIDES}", size=8.5, color=GREY, align=PP_ALIGN.RIGHT)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_table(slide, left, top, width, height, data, widths, *, font_size=11, header_size=11, row_height=None, align_from=1):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
    table = shape.table
    for index, col_width in enumerate(widths):
        table.columns[index].width = Inches(col_width)
    if row_height:
        for row in table.rows:
            row.height = Inches(row_height)
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = ""
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = BROWN if r == 0 else (WHITE if r % 2 else PALE)
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.text = str(value)
            paragraph.alignment = PP_ALIGN.LEFT if c < align_from else PP_ALIGN.CENTER
            for run in paragraph.runs:
                run.font.name = FONT
                run.font.size = Pt(header_size if r == 0 else font_size)
                run.font.bold = r == 0
                run.font.color.rgb = WHITE if r == 0 else INK
    return table


def add_image_fit(slide, path, left, top, width, height):
    from PIL import Image

    with Image.open(path) as image:
        image_w, image_h = image.size
    ratio = image_w / image_h
    if ratio >= width / height:
        draw_w, draw_h = width, width / ratio
    else:
        draw_h, draw_w = height, height * ratio
    draw_left = left + (width - draw_w) / 2
    draw_top = top + (height - draw_h) / 2
    return slide.shapes.add_picture(str(path), Inches(draw_left), Inches(draw_top), Inches(draw_w), Inches(draw_h))


# ---------------------------------------------------------------------------
# 1. Copertina
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide, BROWN_DARK)
add_rect(slide, 0, 0, 13.333, 0.09, fill=GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
add_rect(slide, 0.75, 1.42, 0.9, 0.06, fill=GOLD, line=None, shape=MSO_SHAPE.RECTANGLE)
add_text(slide, 0.75, 1.0, 11.5, 0.3, "PROGETTO FINALE · LLM & AGENTIC AI", size=11, bold=True, color=GOLD)
add_text(slide, 0.75, 1.7, 11.8, 1.5, "Casa & Regali", size=52, bold=True, color=CREAM)
add_text(slide, 0.75, 2.95, 11.2, 0.8, "Assistente clienti in italiano", size=26, color=GOLD)
add_text(
    slide, 0.75, 3.85, 9.8, 1.0,
    "L'LLM sceglie il tool e scrive la risposta.\nIl codice impone policy, privacy, HITL e ledger.",
    size=17, color=CREAM, line_spacing=1.12,
)
chips = ["LangChain", "LangGraph", "BM25", "HITL", "Streamlit"]
x = 0.75
for chip in chips:
    width = 0.42 + 0.115 * len(chip)
    add_chip(slide, x, 5.25, width, 0.44, chip, fill=RGBColor(0x5A, 0x43, 0x32), color=CREAM, size=11.5)
    x += width + 0.18
add_text(slide, 0.75, 6.35, 8.0, 0.3, "LFM3-1.7B Q4_K_M · LM Studio · 29 ottobre 2026", size=12, color=RGBColor(0xD8, 0xCB, 0xBC))
add_notes(slide, "Titolo e tesi: la v2 tiene solo le due decisioni affidate all'LLM e sposta i vincoli nel codice.")

# ---------------------------------------------------------------------------
# 2. In una frase
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "La tesi", "In una frase", "Tre passaggi, un solo agente, nessun fallback deterministico.")
cards = [
    ("1", "L'LLM sceglie", "Riceve la richiesta e la conversazione, e seleziona uno dei sette tool con i suoi argomenti. Se non serve un tool, risponde direttamente."),
    ("2", "Il tool esegue", "Legge i dati del cliente dal runtime e applica la policy: finestra resi, categorie escluse, soglie, privacy. Il modello non può cambiare cliente."),
    ("3", "L'LLM racconta", "Un secondo passo, senza tool, riceve gli esiti e scrive la risposta finale. Non può avviare altre azioni."),
]
for index, (number, title, body) in enumerate(cards):
    left = 0.6 + index * 4.12
    add_rect(slide, left, 1.95, 3.85, 3.0, fill=WHITE)
    add_text(slide, left + 0.25, 2.12, 0.8, 0.6, number, size=30, bold=True, color=GOLD)
    add_text(slide, left + 0.25, 2.78, 3.35, 0.4, title, size=18, bold=True, color=BROWN_DARK)
    add_text(slide, left + 0.25, 3.35, 3.35, 1.45, body, size=12.5, color=GREY, line_spacing=1.06)
add_rect(slide, 0.6, 5.4, 12.13, 0.85, fill=LIGHT_BROWN, line=None)
add_text(
    slide, 0.85, 5.57, 11.6, 0.55,
    "Il rimborso è solo una proposta: interrupt, approvazione umana, poi ledger idempotente e audit JSONL.",
    size=14, bold=True, color=BROWN_DARK, anchor=MSO_ANCHOR.MIDDLE,
)
add_footer(slide, 2)
add_notes(slide, "Sottolineare la separazione tra decisione (LLM) ed esecuzione vincolata (tool + policy).")

# ---------------------------------------------------------------------------
# 3. Da estesa a v2
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "La revisione", "Da estesa a v2", "Meno pattern, stessi vincoli di dominio.")
add_card(
    slide, 0.6, 1.75, 5.15, 3.0, "Versione estesa",
    "", accent=RGBColor(0x8A, 0x7A, 0x66), body_size=13,
    body_lines=[
        "Supervisore + 3 agenti",
        "Ciclo ReAct con più passaggi",
        "Review / reflection",
        "Chroma + embedding",
        "Backend deterministico",
        "HITL come nodo separato",
    ],
)
add_card(
    slide, 7.55, 1.75, 5.15, 3.0, "Versione 2 (v2)",
    "", accent=GREEN, fill=LIGHT_GREEN, body_size=13,
    body_lines=[
        "1 agente",
        "Un solo passaggio di tool",
        "Nessuna reflection",
        "BM25 su 10 documenti",
        "Nessun fallback",
        "HITL dentro proponi_rimborso",
    ],
)
arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.92), Inches(3.0), Inches(1.45), Inches(0.5))
arrow.fill.solid()
arrow.fill.fore_color.rgb = GOLD
arrow.line.fill.background()
arrow.shadow.inherit = False
add_text(slide, 0.6, 5.15, 12.13, 1.0,
         "La versione estesa dimostra più pattern, ma con un modello da 1,7B il routing e la reflection aggiungono latenza e punti di errore. "
         "La revisione toglie ciò che non ripaga e conserva policy, privacy, HITL, audit e ledger.",
         size=13.5, color=GREY, line_spacing=1.08)
add_footer(slide, 3)
add_notes(slide, "Non è una critica alla versione estesa: è una scelta proporzionata al modello locale e alla leggibilità.")

# ---------------------------------------------------------------------------
# 4. Architettura
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Architettura", "Quattro nodi, un solo passaggio di tool", "agent → tools → respond → finalize; senza tool: agent → finalize.")
add_image_fit(slide, ARCH, 0.6, 1.72, 12.13, 4.1)
add_chip(slide, 0.62, 6.0, 3.9, 0.5, "agent → tools → respond → finalize", fill=LIGHT_BROWN, size=11.5)
add_chip(slide, 4.72, 6.0, 3.4, 0.5, "agent → finalize (senza tool)", fill=LIGHT_GREEN, color=GREEN, size=11.5)
add_chip(slide, 8.32, 6.0, 4.4, 0.5, "dopo il rimborso: nessun altro tool", fill=LIGHT_RED, color=RGBColor(0x8A, 0x4B, 0x3F), size=11.5)
add_footer(slide, 4)
add_notes(slide, "Il secondo passo LLM non ha tool: scrive solo la risposta. Dopo proponi_rimborso il turno è chiuso.")

# ---------------------------------------------------------------------------
# 5. Cosa fa chi
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Confine", "Cosa fa l'LLM, cosa fa il codice", "Due decisioni al modello, tutti i vincoli al codice.")
add_card(
    slide, 0.6, 1.75, 5.85, 3.05, "LLM", "", accent=GOLD, body_size=14,
    body_lines=[
        "Sceglie il tool e i suoi argomenti",
        "Scrive la risposta in linguaggio naturale",
        "Può sbagliare: l'errore resta visibile nella valutazione",
        "Non può cambiare cliente, importo o esito",
    ],
)
add_card(
    slide, 6.88, 1.75, 5.85, 3.05, "Codice", "", accent=GREEN, fill=LIGHT_GREEN, body_size=14,
    body_lines=[
        "policy.py applica finestra resi, esclusioni e soglie",
        "ToolRuntime impone il cliente della sessione",
        "interrupt() sospende il rimborso per l'approvazione",
        "audit.py scrive solo dopo approvazione e registra ogni turno",
    ],
)
add_rect(slide, 0.6, 5.2, 12.13, 0.62, fill=PALE, line=None)
add_text(slide, 0.85, 5.32, 11.6, 0.4,
         "Nessuna modalità deterministica, nessun fallback semantico: se il modello sbaglia, lo si vede.",
         size=13.5, bold=True, color=BROWN_DARK)
add_footer(slide, 5)
add_notes(slide, "Il punto centrale: l'AI è responsabile delle decisioni, il codice garantisce che non possa violare i vincoli.")

# ---------------------------------------------------------------------------
# 6. I tool
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Strumenti", "Sette tool, quattro read-only", "Solo proponi_rimborso scrive un rimborso, e solo dopo l'approvazione umana.")
add_table(
    slide, 0.6, 1.68, 12.13, 4.55,
    [
        ["Tool", "Cosa fa", "Scrittura"],
        ["cerca_policy(query)", "Cerca nelle policy e restituisce contesto e fonti", "no"],
        ["stato_ordine(order_id)", "Stato, articoli e tracking di un ordine del cliente", "no"],
        ["elenca_ordini()", "Elenco degli ordini del profilo della sessione", "no"],
        ["verifica_reso(order_id)", "Idoneità al reso, giorni trascorsi, esclusioni", "no"],
        ["proponi_rimborso(order_id)", "Verifica, chiede HITL, scrive solo se approvato", "dopo HITL"],
        ["apri_ticket(motivo)", "Apre un ticket per un operatore umano", "ticket"],
        ["ricorda_preferenza(testo)", "Salva una preferenza su richiesta esplicita", "preferenza"],
    ],
    [3.55, 6.6, 1.98],
    font_size=11.5,
    header_size=11.5,
    row_height=0.5,
    align_from=2,
)
add_text(slide, 0.6, 5.85, 12.1, 0.4,
         "Tutti restituiscono testo leggibile, tranne cerca_policy che restituisce JSON con contesto e fonti.",
         size=11.5, italic=True, color=GREY)
add_footer(slide, 6)
add_notes(slide, "I tool leggono i dati dal runtime: il modello non può indicare un cliente diverso da quello della sessione.")

# ---------------------------------------------------------------------------
# 7. HITL
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Rimborso", "Human-in-the-Loop dentro il tool", "La proposta è visibile, l'esecuzione no: il ledger scrive solo dopo l'ok umano.")
steps = [
    ("1. Verifica", "idoneità, finestra, esclusioni, duplicati"),
    ("2. Proposta", "ordine, cliente, importo, livello"),
    ("3. Interrupt", "il turno si sospende"),
    ("4. Decisione", "approva o rifiuta in UI o CLI"),
    ("5. Ledger", "rivalida e scrive solo se approvato"),
]
for index, (title, body) in enumerate(steps):
    left = 0.62 + index * 2.46
    add_rect(slide, left, 1.95, 2.2, 1.5, fill=LIGHT_RED if index == 2 else WHITE)
    add_text(slide, left + 0.16, 2.08, 1.9, 0.35, title, size=14, bold=True, color=BROWN_DARK)
    add_text(slide, left + 0.16, 2.48, 1.88, 0.85, body, size=11, color=GREY, line_spacing=1.02)
    if index < 4:
        add_arrow(slide, left + 2.23, 2.6, width=0.2, color=GOLD)
add_card(
    slide, 0.6, 3.95, 5.85, 1.55, "Il ledger non si fida",
    "Un importo manomesso, un'approvazione mancante o un rimborso già registrato vengono rifiutati. "
    "Il chiamante non può forzare la scrittura.",
    accent=RGBColor(0x8A, 0x4B, 0x3F), body_size=12.5,
)
add_card(
    slide, 6.88, 3.95, 5.85, 1.55, "Soglia 500 EUR",
    "Sopra i 500 EUR la proposta richiede il livello responsabile: la UI mostra una conferma aggiuntiva, "
    "la CLI la chiede esplicitamente.",
    accent=GOLD, body_size=12.5,
)
add_rect(slide, 0.6, 5.85, 12.13, 0.68, fill=LIGHT_BROWN, line=None)
add_text(slide, 0.85, 5.99, 11.6, 0.4,
         "Senza approved=True non esiste alcun rimborso: il ledger rivalida cliente, importo e duplicati.",
         size=13.5, bold=True, color=BROWN_DARK)
add_footer(slide, 7)
add_notes(slide, "HITL dentro proponi_rimborso tramite interrupt(): nessun nodo separato, nessuna scrittura prima dell'ok.")

# ---------------------------------------------------------------------------
# 8. Dati, RAG, memoria
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Dati e contesto", "Dati, RAG e memoria", "Un corpus piccolo, un indice lessicale, una memoria esplicita.")
add_card(
    slide, 0.6, 1.75, 3.85, 3.15, "Dati demo", "", accent=BROWN, body_size=12.5,
    body_lines=[
        "20 prodotti, 50 clienti",
        "200 ordini, 4 rimborsi iniziali",
        "Casi curati: ORD-1001…1008",
        "Nessun dato personale reale",
        "DEMO_TODAY = 2026-10-23",
    ],
)
add_card(
    slide, 4.74, 1.75, 3.85, 3.15, "RAG BM25", "", accent=GREEN, fill=LIGHT_GREEN, body_size=12.5,
    body_lines=[
        "10 documenti Markdown",
        "chunk_size = 650",
        "top_k = 3",
        "Niente embedding o Chroma",
        "Query + ultima richiesta del cliente",
    ],
)
add_card(
    slide, 8.88, 1.75, 3.85, 3.15, "Memoria", "", accent=GOLD, body_size=12.5,
    body_lines=[
        "Breve: InMemorySaver",
        "Lunga: PreferenceMemory",
        "Solo preferenze richieste esplicitamente",
        "Email, telefoni e carte rimossi",
        "Stato di sessione perso al riavvio",
    ],
)
add_rect(slide, 0.6, 5.25, 12.13, 0.7, fill=PALE, line=None)
add_text(slide, 0.85, 5.39, 11.6, 0.42,
         "La query del modello viene unita all'ultima richiesta: i termini specifici del cliente non si perdono.",
         size=13, bold=True, color=BROWN_DARK)
add_footer(slide, 8)
add_notes(slide, "BM25 è proporzionato a dieci documenti brevi. La memoria lunga non salva nulla senza richiesta esplicita.")

# ---------------------------------------------------------------------------
# 9. Valutazione
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Valutazione", "31 scenari, due profili", "Tool, contenuti, fonti, HITL, scritture ed effetti collaterali.")
stats = [
    (f"{fs['passed']}/{fs['total']}", "LFM3 superati"),
    (fmt_pct(fs["tool_hit_rate"]), "tool hit LFM3"),
    (f"{ls.get('passed', 'n/d')}/{ls.get('total', 31)}", "Llama superati"),
    (fmt_s(ls.get("avg_latency_s")), "latenza Llama"),
]
for index, (value, label) in enumerate(stats):
    left = 0.6 + index * 3.06
    add_rect(slide, left, 1.62, 2.9, 0.95, fill=PALE if index % 2 == 0 else LIGHT_GREEN, line=None)
    add_text(slide, left + 0.12, 1.69, 2.66, 0.4, value, size=20, bold=True, color=BROWN, align=PP_ALIGN.CENTER)
    add_text(slide, left + 0.12, 2.13, 2.66, 0.25, label, size=10, color=GREY, align=PP_ALIGN.CENTER)
add_image_fit(slide, CHART, 0.85, 2.72, 11.6, 2.75)
add_text(slide, 0.6, 5.62, 12.1, 0.9,
         f"LFM3: {fmt_s(fs['avg_latency_s'])} per turno, 100% su tool, fonti, HITL e sicurezza. "
         f"Llama: {fmt_s(ls.get('avg_latency_s'))} ma {fmt_pct(ls.get('safety_pass_rate'))} di sicurezza e {fmt_pct(ls.get('source_hit_rate'))} sulle fonti.",
         size=12.5, color=GREY, line_spacing=1.06)
add_text(slide, 0.6, 6.45, 12.1, 0.35,
         "31 scenari: 9 info, 3 ordini, 6 resi, 6 rimborsi, 3 sicurezza, 2 escalation, 2 memoria.",
         size=11, italic=True, color=GREY)
add_footer(slide, 9)
add_notes(slide, "Numeri da eval/results/latest.json e dalla run Llama completa. Gemini non ha una run comparabile per quota.")

# ---------------------------------------------------------------------------
# 10. Errori e reasoning
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Analisi", "Dove si concentrano gli errori", "Llama è veloce ma perde resi, memoria e fonti.")
add_card(
    slide, 0.6, 1.75, 3.85, 3.15, "Llama 3.2 3B (Unsloth)", "", accent=RGBColor(0x8A, 0x4B, 0x3F), body_size=12.5,
    body_lines=[
        "Informazioni 8/9",
        "Resi 3/6",
        "Privacy/sicurezza 2/3",
        "Memoria 0/2",
        "Totale 24/31",
    ],
)
add_card(
    slide, 4.74, 1.75, 3.85, 3.15, "Reasoning LFM3", "", accent=GREEN, fill=LIGHT_GREEN, body_size=12.5,
    body_lines=[
        "Tool: reasoning default, 31/31",
        "Risposta finale: reasoning off",
        "Prima: 7.058 token, output vuoto",
        "none su tutto il grafo: 27/31",
    ],
)
add_card(
    slide, 8.88, 1.75, 3.85, 3.15, "Gemini 2.5 Flash", "", accent=RGBColor(0x4D, 0x5B, 0x73), fill=LIGHT_BLUE, body_size=12.5,
    body_lines=[
        "Backend configurabile",
        "Quota gratuita giornaliera esaurita",
        "Nessuna run completa",
        "Nessuna metrica comparabile inventata",
    ],
)
add_rect(slide, 0.6, 5.25, 12.13, 0.65, fill=LIGHT_BROWN, line=None)
add_text(slide, 0.85, 5.37, 11.6, 0.42,
         "La valutazione misura gli errori, non li nasconde: è il motivo per cui il fallback deterministico è stato escluso.",
         size=13.5, bold=True, color=BROWN_DARK)
add_footer(slide, 10)
add_notes(slide, "Il reasoning resta utile per scegliere il tool; la risposta finale lo disattiva per evitare l'esaurimento del contesto. Disattivarlo su tutto il grafo peggiora i risultati.")

# ---------------------------------------------------------------------------
# 11. Test e riproducibilità
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Qualità", "Test e riproducibilità", "La suite automatica non richiede LM Studio; la valutazione sì.")
add_card(
    slide, 0.6, 1.75, 4.0, 3.3, "45 test", "", accent=GREEN, fill=LIGHT_GREEN, body_size=12.5,
    body_lines=[
        "ScriptedModel, nessuna rete",
        "Grafo, HITL, privacy, memoria",
        "Policy, corpus, ledger, scenari",
        "Dati isolati per ogni test",
    ],
)
add_rect(slide, 4.9, 1.75, 7.83, 3.3, fill=WHITE)
add_text(slide, 5.12, 1.95, 7.4, 0.3, "Comandi", size=15, bold=True, color=BROWN_DARK)
add_text(
    slide, 5.12, 2.4, 7.35, 2.45,
    "python -m pytest -q tests\n\npython eval/run_eval.py --backend local\n\nstreamlit run app.py\n\npython cli.py --customer C-001",
    size=11.5, color=INK, font=FONT_MONO, line_spacing=1.15,
)
add_rect(slide, 0.6, 5.3, 12.13, 0.78, fill=PALE, line=None)
add_text(
    slide, 0.85, 5.4, 11.6, 0.58,
    "LM Studio: lms load lfm3-1.7b@q4_k_m --gpu max --context-length 8192 --yes\n"
    "Reset demo: python scripts/reset_demo.py · risultati in eval/results/",
    size=11.5, color=GREY, font=FONT_MONO, line_spacing=1.1,
)
add_footer(slide, 11)
add_notes(slide, "I test usano un modello scriptato e non richiedono LM Studio. La run reale richiede il modello caricato.")

# ---------------------------------------------------------------------------
# 12. Limiti e conclusione
# ---------------------------------------------------------------------------
slide = prs.slides.add_slide(BLANK)
set_bg(slide)
add_title_block(slide, "Chiusura", "Limiti e conclusione", "Un prototipo misurato, non una garanzia di produzione.")
add_bullets(
    slide, 0.6, 1.8, 7.5, 4.6,
    [
        "Modelli piccoli: possono sbagliare tool o riformulare i numeri; la valutazione lo mostra.",
        "Un tool per turno: una richiesta che ne richiede due viene gestita al turno successivo.",
        "BM25: adatto a dieci documenti brevi, non a un corpus grande.",
        "Memoria di sessione: InMemorySaver perde lo stato al riavvio.",
        "Gemini: backend presente, run completa non disponibile per quota.",
    ],
    size=13.5,
    gap=12,
)
add_rect(slide, 8.45, 1.8, 4.28, 3.4, fill=LIGHT_BROWN, line=GOLD)
add_text(slide, 8.7, 2.0, 3.8, 0.4, "Il punto", size=18, bold=True, color=BROWN_DARK)
add_text(
    slide, 8.7, 2.55, 3.75, 2.45,
    "Il grafo è piccolo abbastanza da essere letto in un minuto, ma il comportamento non è simulato. "
    "L'LLM sceglie davvero, scrive davvero e sbaglia quando il modello sbaglia. La differenza è che "
    "gli errori sono visibili, non nascosti da un fallback.",
    size=12.5, color=GREY, line_spacing=1.08,
)
add_text(
    slide, 0.6, 5.6, 12.1, 0.45,
    "31/31 su 31 scenari controllati, con modello locale e dati sintetici: una misura, non una promessa.",
    size=13, italic=True, color=BROWN, align=PP_ALIGN.CENTER,
)
add_footer(slide, 12)
add_notes(slide, "Chiudere sul confine tra AI e codice: è il contributo principale della v2.")

DOCS.mkdir(exist_ok=True)
prs.save(OUT)
print(f"Scritto {OUT}")
